# -*- coding: utf-8 -*-
"""
Import title-engine keywords from a manually prepared CSV/XLSX file.

The command is dry-run by default. Use --write-db to upsert rows into the
formal JSReportReplica/app title keyword table.
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path
from typing import Any, Iterable


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))


COLUMN_ALIASES = {
    "keyword": ("keyword", "word", "关键词", "搜索词", "标题词"),
    "category": ("category", "类目", "分类", "商品类目"),
    "category_alias": ("category_alias", "类目别名", "分类别名"),
    "source": ("source", "来源", "数据来源"),
    "search_popularity": ("search_popularity", "搜索人气", "搜索热度"),
    "transaction_index": ("transaction_index", "交易指数", "成交指数"),
    "competition": ("competition", "竞争度", "竞争指数"),
    "relevance": ("relevance", "相关性", "相关度"),
    "product_fit": ("product_fit", "产品适配度", "商品适配度", "适配度"),
    "source_stat_date": ("source_stat_date", "统计日期", "数据日期"),
    "source_record_key": ("source_record_key", "来源记录键", "源记录键"),
}


def value_for(row: dict[str, Any], field_name: str, default: Any = "") -> Any:
    """Read a field value using configured source column aliases."""
    for alias in COLUMN_ALIASES[field_name]:
        if alias in row and row[alias] not in (None, ""):
            return row[alias]
    return default


def parse_number(value: Any) -> float:
    """Parse a human-entered numeric cell."""
    text = str(value or "").strip().replace(",", "")
    if text.endswith("%"):
        text = text[:-1]
    if not text:
        return 0.0
    try:
        return float(text)
    except ValueError:
        return 0.0


def parse_optional_number(value: Any) -> float | None:
    """Parse a score input while preserving missing or invalid cells."""
    text = str(value or "").strip().replace(",", "")
    if text.endswith("%"):
        text = text[:-1]
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def normalize_keyword_row(row: dict[str, Any], *, default_source: str = "manual") -> dict[str, Any]:
    """Normalize one CSV/XLSX row into the title keyword contract."""
    cleaned = {str(key or "").strip(): value for key, value in row.items() if str(key or "").strip()}
    keyword = str(value_for(cleaned, "keyword")).strip()
    category = str(value_for(cleaned, "category")).strip()
    if not keyword:
        raise ValueError("keyword is required")
    if not category:
        raise ValueError("category is required")

    source = str(value_for(cleaned, "source", default_source) or default_source).strip() or default_source
    relevance = parse_optional_number(value_for(cleaned, "relevance", None))
    raw_product_fit = value_for(cleaned, "product_fit", "")
    product_fit = parse_optional_number(raw_product_fit) if raw_product_fit not in (None, "") else relevance
    return {
        "keyword": keyword,
        "category": category,
        "category_alias": str(value_for(cleaned, "category_alias", "") or "").strip(),
        "source": source,
        "search_popularity": parse_optional_number(value_for(cleaned, "search_popularity", None)),
        "transaction_index": parse_number(value_for(cleaned, "transaction_index", 0)),
        "competition": parse_optional_number(value_for(cleaned, "competition", None)),
        "relevance": relevance,
        "product_fit": product_fit,
        "source_stat_date": str(value_for(cleaned, "source_stat_date", "") or "").strip(),
        "source_record_key": str(value_for(cleaned, "source_record_key", "") or "").strip(),
        "raw_payload": json.dumps(cleaned, ensure_ascii=False, default=str),
    }


def read_csv_rows(path: Path) -> list[dict[str, Any]]:
    """Read rows from a UTF-8/UTF-8-BOM CSV file."""
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [dict(row) for row in csv.DictReader(handle)]


def read_xlsx_rows(path: Path, *, sheet_name: str = "") -> list[dict[str, Any]]:
    """Read rows from an XLSX/XLSM workbook."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for Excel import") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value or "").strip() for value in rows[0]]
    result: list[dict[str, Any]] = []
    for values in rows[1:]:
        result.append({headers[index]: value for index, value in enumerate(values) if index < len(headers)})
    return result


def load_keyword_rows(path: Path, *, sheet_name: str = "", default_source: str = "manual") -> list[dict[str, Any]]:
    """Load and normalize keyword rows from CSV/XLSX."""
    suffix = path.suffix.lower()
    if suffix == ".csv":
        raw_rows = read_csv_rows(path)
    elif suffix in {".xlsx", ".xlsm"}:
        raw_rows = read_xlsx_rows(path, sheet_name=sheet_name)
    else:
        raise ValueError(f"Unsupported keyword file type: {suffix}")

    normalized_rows: list[dict[str, Any]] = []
    errors: list[str] = []
    for index, row in enumerate(raw_rows, start=2):
        try:
            normalized_rows.append(normalize_keyword_row(row, default_source=default_source))
        except ValueError as exc:
            errors.append(f"row {index}: {exc}")
    if errors:
        raise ValueError("; ".join(errors[:10]))
    return normalized_rows


def iter_limited(rows: list[dict[str, Any]], limit: int) -> Iterable[dict[str, Any]]:
    """Iterate over the first N rows when a positive limit is set."""
    return rows[:limit] if limit > 0 else rows


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Import manual title keywords from CSV/XLSX.")
    parser.add_argument("--file", required=True, help="CSV/XLSX keyword file.")
    parser.add_argument("--sheet", default="", help="Excel sheet name. Defaults to the active sheet.")
    parser.add_argument("--source", default="manual", help="Keyword source label.")
    parser.add_argument("--updated-by", default="manual_import", help="Operator recorded in created_by.")
    parser.add_argument("--limit", type=int, default=0, help="Limit rows for a controlled import.")
    parser.add_argument("--write-db", action="store_true", help="Actually upsert rows into JSReportReplica/app.")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    path = Path(args.file).resolve()
    rows = list(iter_limited(load_keyword_rows(path, sheet_name=args.sheet, default_source=args.source), args.limit))
    summary = {
        "file": str(path),
        "sheet": args.sheet,
        "source": args.source,
        "loaded_count": len(rows),
        "write_db": bool(args.write_db),
        "upserted_count": 0,
        "preview": rows[:5],
    }

    if args.write_db:
        from database.db_manager import DatabaseManager

        manager = DatabaseManager()
        for row in rows:
            if manager.upsert_keyword(row, updated_by=args.updated_by):
                summary["upserted_count"] += 1

    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
